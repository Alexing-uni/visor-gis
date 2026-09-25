"""Rebuild docs/Visor_GIS_0.4.docx from docs/MEMORIA.md using python-docx.

Run with a Python environment that includes python-docx:
  python scripts/generate-docx.py
Only standard-library modules and python-docx are required. This program creates
native headings, multilevel numbering, tables, bookmarks, links and a TOC field.
The TOC has cached linked entries without guessed page numbers; Word refreshes it.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
BLACK = RGBColor(0, 0, 0)


def node(name: str, **attributes):
    result = OxmlElement(name)
    for key, value in attributes.items():
        result.set(qn(key), str(value))
    return result


def hyperlink(paragraph, label: str, target: str, internal: bool = False):
    link = node('w:hyperlink')
    if internal:
        link.set(qn('w:anchor'), target)
    else:
        rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
        link.set(qn('r:id'), rel)
    run = node('w:r')
    properties = node('w:rPr')
    properties.append(node('w:color', **{'w:val': '000000'}))
    if not internal:
        properties.append(node('w:u', **{'w:val': 'single'}))
    run.append(properties)
    text = node('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, text: str):
    # Controlled source: inline code, bold and literal HTTPS links.
    pieces = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|https?://[^\s)]+)', text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif piece.startswith('**') and piece.endswith('**'):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith(('https://', 'http://')):
            cleaned = piece.rstrip('.,;')
            hyperlink(paragraph, cleaned, cleaned)
            paragraph.add_run(piece[len(cleaned):])
        else:
            paragraph.add_run(piece)


def configure(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.footer_distance = Cm(1.1)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = BLACK
        style.font.underline = False
        style.paragraph_format.widow_control = True
        if style._element.pPr is not None:
            for borders in list(style._element.pPr.findall(qn('w:pBdr'))):
                style._element.pPr.remove(borders)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(7)
    title = document.styles['Title']
    title.font.size = Pt(25)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)
    document.styles['Subtitle'].font.size = Pt(14)
    for name, size, before, after in [('Heading 1', 16, 17, 8), ('Heading 2', 12.5, 12, 6), ('Heading 3', 11.5, 10, 5)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    code = document.styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = 'Consolas'
    code.font.size = Pt(9)
    code.font.color.rgb = BLACK
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.widow_control = True
    for level in (1, 2):
        toc = document.styles.add_style(f'Manual TOC {level}', WD_STYLE_TYPE.PARAGRAPH)
        toc.font.name = 'Calibri'
        toc.font.size = Pt(10.5)
        toc.font.color.rgb = BLACK
        toc.paragraph_format.space_after = Pt(4)
        toc.paragraph_format.line_spacing = 1.0
        toc.paragraph_format.left_indent = Cm(0 if level == 1 else 0.55)
        toc.font.bold = level == 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Visor GIS 0.4  |  ')
    run.font.size = Pt(9)
    field = node('w:fldSimple', **{'w:instr': 'PAGE'})
    field.append(node('w:r'))
    field[0].append(node('w:t'))
    field[0][0].text = '1'
    footer._p.append(field)
    document.settings.element.append(node('w:updateFields', **{'w:val': 'true'}))
    document.core_properties.title = 'Visor GIS web versión 0.4'
    document.core_properties.subject = 'Memoria técnica y manual de uso'
    document.core_properties.author = 'Proyecto Visor GIS'
    document.core_properties.keywords = 'GIS, deck.gl, MapLibre, React, análisis, rutas, GitHub Pages'
    document.core_properties.comments = ''


def heading_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_id = max([int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))] + [-1]) + 1
    number_id = max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))] + [0]) + 1
    abstract = node('w:abstractNum', **{'w:abstractNumId': abstract_id})
    abstract.append(node('w:multiLevelType', **{'w:val': 'multilevel'}))
    for level in range(3):
        lvl = node('w:lvl', **{'w:ilvl': level})
        lvl.append(node('w:start', **{'w:val': 1}))
        lvl.append(node('w:numFmt', **{'w:val': 'decimal'}))
        lvl.append(node('w:pStyle', **{'w:val': f'Heading{level + 1}'}))
        lvl.append(node('w:lvlText', **{'w:val': '.'.join(f'%{i + 1}' for i in range(level + 1))}))
        lvl.append(node('w:suff', **{'w:val': 'space'}))
        lvl.append(node('w:lvlJc', **{'w:val': 'left'}))
        abstract.append(lvl)
    numbering.append(abstract)
    number = node('w:num', **{'w:numId': number_id})
    number.append(node('w:abstractNumId', **{'w:val': abstract_id}))
    numbering.append(number)
    return number_id


def add_heading(document, text: str, level: int, number_id: int, bookmark_id: int):
    match = re.match(r'([\d.]+)\s+(.+)', text)
    prefix, clean = (match[1], match[2]) if match else ('', text)
    paragraph = document.add_heading(clean, level)
    properties = paragraph._p.get_or_add_pPr()
    num = node('w:numPr')
    num.append(node('w:ilvl', **{'w:val': level - 1}))
    num.append(node('w:numId', **{'w:val': number_id}))
    properties.append(num)
    name = 'sec_' + prefix.replace('.', '_') if prefix else f'sec_{bookmark_id}'
    start = node('w:bookmarkStart', **{'w:id': bookmark_id, 'w:name': name})
    end = node('w:bookmarkEnd', **{'w:id': bookmark_id})
    paragraph._p.insert(1, start)
    paragraph._p.append(end)
    return name


def add_toc(document, headings):
    paragraph = document.add_paragraph('Índice')
    paragraph.runs[0].bold = True
    paragraph.runs[0].font.size = Pt(16)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.keep_with_next = True
    first = document.add_paragraph(style='Manual TOC 1')
    first.add_run()._r.append(node('w:fldChar', **{'w:fldCharType': 'begin', 'w:dirty': 'true'}))
    instruction = node('w:instrText', **{'xml:space': 'preserve'})
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    first.add_run()._r.append(instruction)
    first.add_run()._r.append(node('w:fldChar', **{'w:fldCharType': 'separate'}))
    current = first
    for index, (level, heading) in enumerate(headings):
        if index:
            current = document.add_paragraph(style=f'Manual TOC {level}')
        else:
            current.style = document.styles[f'Manual TOC {level}']
        prefix = re.match(r'([\d.]+)', heading)[1]
        hyperlink(current, heading, 'sec_' + prefix.replace('.', '_'), internal=True)
    current.add_run()._r.append(node('w:fldChar', **{'w:fldCharType': 'end'}))
    note = document.add_paragraph('El índice enlaza con los apartados. En Word, haz clic derecho sobre él y elige Actualizar campo y Actualizar toda la tabla para incorporar la paginación. Los títulos conservan sus estilos y numeración automática.')
    for run in note.runs:
        run.font.size = Pt(9.5)


def add_table(document, rows):
    columns = len(rows[0])
    table = document.add_table(rows=0, cols=columns)
    table.autofit = False
    widths = {2: [5.5, 11.1], 3: [4.2, 4.0, 8.4], 5: [5.2, 1.8, 3.3, 3.0, 3.3]}.get(columns, [16.6 / columns] * columns)
    props = table._tbl.tblPr
    borders = node('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(node(f'w:{edge}', **{'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
    props.append(borders)
    margins = node('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        margins.append(node(f'w:{edge}', **{'w:w': 85, 'w:type': 'dxa'}))
    props.append(margins)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        tr_props = row._tr.get_or_add_trPr()
        tr_props.append(node('w:cantSplit'))
        if index == 0:
            tr_props.append(node('w:tblHeader', **{'w:val': 'true'}))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = 'D9E5EF' if index == 0 else ('F6F8FA' if index % 2 == 0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(node('w:shd', **{'w:fill': shading}))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            if columns == 5 and j in [1, 3, 4]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = BLACK
                if index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build(source: Path, output: Path):
    text = source.read_text(encoding='utf-8')
    lines = text.splitlines()
    headings = [(len(match[1]) - 1, match[2]) for line in lines if (match := re.match(r'^(#{2,3})\s+(.+)$', line))]
    document = Document()
    configure(document)
    number_id = heading_numbering(document)
    first_section = next(i for i, line in enumerate(lines) if line.startswith('## '))
    document.add_paragraph(lines[0].removeprefix('# '), 'Title')
    for line in lines[1:first_section]:
        if line.strip():
            if line == 'Memoria técnica y manual de uso':
                document.add_paragraph(line, 'Subtitle')
            else:
                inline(document.add_paragraph(), line)
    document.add_page_break()
    add_toc(document, headings)
    document.add_page_break()
    index = first_section
    bookmark_id = 1
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        heading = re.match(r'^(#{2,4})\s+(.+)$', line)
        if heading:
            add_heading(document, heading[2], len(heading[1]) - 1, number_id, bookmark_id)
            bookmark_id += 1
            index += 1
        elif line.startswith('```'):
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith('```'):
                block.append(lines[index])
                index += 1
            for position, code_line in enumerate(block):
                paragraph = document.add_paragraph(style='Code Block')
                paragraph.add_run(code_line)
                paragraph.paragraph_format.keep_with_next = position < len(block) - 1 and len(block) < 12
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            index += 1
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                row = [part.strip() for part in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', item) for item in row):
                    rows.append(row)
                index += 1
            add_table(document, rows)
        elif line.startswith('- '):
            inline(document.add_paragraph(style='List Bullet'), line[2:])
            index += 1
        else:
            paragraph_lines = [line]
            index += 1
            while index < len(lines) and lines[index].strip() and not lines[index].startswith(('#', '```', '|', '- ')):
                paragraph_lines.append(lines[index])
                index += 1
            inline(document.add_paragraph(), ' '.join(paragraph_lines))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    verify = Document(output)
    sections = [p.text for p in verify.paragraphs if p.style.name == 'Heading 1']
    if len(sections) != 15:
        raise RuntimeError(f'Expected 15 main sections, found {len(sections)}')
    if len(verify.tables) != 4:
        raise RuntimeError(f'Expected 4 tables, found {len(verify.tables)}')
    print(f'Created {output}; {len(sections)} numbered sections, {len(verify.tables)} tables; TOC, bookmarks and PAGE field present.')
    print('Structural verification passed. Visual verification requires rendering the DOCX in a compatible office renderer.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=ROOT / 'docs' / 'MEMORIA.md')
    parser.add_argument('--output', type=Path, default=ROOT / 'docs' / 'Visor_GIS_0.4.docx')
    arguments = parser.parse_args()
    build(arguments.source, arguments.output)
